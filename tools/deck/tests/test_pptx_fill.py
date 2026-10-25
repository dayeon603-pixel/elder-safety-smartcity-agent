import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Pt

import pptx_fill as pf


@pytest.fixture()
def template(tmp_path: Path) -> Path:
    path = tmp_path / "t.pptx"
    Presentation().save(str(path))
    return path


@pytest.fixture()
def wide_img(tmp_path: Path) -> Path:
    path = tmp_path / "wide.png"
    Image.new("RGB", (400, 100), "red").save(path)
    return path


def write_spec(tmp_path: Path, slides: list[dict]) -> Path:
    path = tmp_path / "spec.json"
    path.write_text(json.dumps({"slides": slides}), encoding="utf-8")
    return path


def test_contain_fit_no_distortion_and_centered() -> None:
    box = pf.Box(100, 200, 1000, 1000)
    fit = pf.contain_fit(400, 100, box)
    assert fit.width == 1000 and fit.height == 250
    assert fit.left == 100 and fit.top == 200 + (1000 - 250) // 2
    tall = pf.contain_fit(100, 400, box)
    assert tall.height == 1000 and tall.width == 250 and tall.left == 100 + 375


def test_contain_fit_invalid() -> None:
    with pytest.raises(pf.SpecError):
        pf.contain_fit(0, 10, pf.Box(0, 0, 10, 10))


def test_inspect_lists_layouts(template: Path) -> None:
    text = "\n".join(pf.inspect_template(template))
    assert "'Title Only'" in text and "type=PICTURE" in text and "idx=1" in text


def test_build_fills_everything(tmp_path: Path, template: Path, wide_img: Path) -> None:
    spec = write_spec(tmp_path, [
        {"layout": "Title and Content", "title": "제목", "body": ["하나", "둘"], "image": str(wide_img),
         "credit": "사진: A, CC0, Wikimedia Commons", "notes": "대본 문장."},
        {"layout": 5, "title": "T2"},
    ])
    out = pf.build_deck(template, spec, tmp_path / "o.pptx")
    prs = Presentation(str(out))
    s1, s2 = prs.slides
    assert s1.shapes.title.text_frame.text == "제목"
    assert s1.notes_slide.notes_text_frame.text == "대본 문장."
    pic = next(sh for sh in s1.shapes if sh.shape_type == 13)
    assert pic.width / pic.height == pytest.approx(4.0, rel=0.01)
    credit = next(sh for sh in s1.shapes if sh.name == "Credit")
    run = credit.text_frame.paragraphs[0].runs[0]
    assert run.font.size == Pt(9) and str(run.font.color.rgb) == "808080"
    assert credit.top > pic.top + pic.height
    assert pic.left >= 0 and pic.left + pic.width <= prs.slide_width
    assert [sh.name for sh in s2.shapes] == ["Title 1"]  # empty placeholders removed


def test_picture_placeholder_replaced_by_contained_image(tmp_path: Path, template: Path, wide_img: Path) -> None:
    spec = write_spec(tmp_path, [{"layout": "Picture with Caption", "title": "t", "image": str(wide_img)}])
    prs = Presentation(str(pf.build_deck(template, spec, tmp_path / "o.pptx")))
    pic = next(sh for sh in prs.slides[0].shapes if sh.shape_type == 13)
    assert pic.width / pic.height == pytest.approx(4.0, rel=0.01)


def test_spec_validation(tmp_path: Path, template: Path) -> None:
    for slides in (
        [{"layout": "Title Only", "body": ["a", "b", "c", "d"]}],
        [{"title": "no layout"}],
        [{"layout": "Nope"}],
        [{"layout": 99}],
        [{"layout": "Title Only", "body": ["x"]}],  # no body placeholder
        [{"layout": "Title Only", "image": "missing.png"}],
    ):
        with pytest.raises(pf.SpecError):
            pf.build_deck(template, write_spec(tmp_path, slides), tmp_path / "o.pptx")


def test_empty_spec(tmp_path: Path) -> None:
    p = tmp_path / "s.json"
    p.write_text('{"slides": []}', encoding="utf-8")
    with pytest.raises(pf.SpecError):
        pf.load_spec(p)


def test_main_returns_1_on_error(tmp_path: Path, template: Path) -> None:
    spec = write_spec(tmp_path, [{"layout": "Nope"}])
    assert pf.main(["build", str(template), str(spec), str(tmp_path / "o.pptx"), "--no-render"]) == 1


def test_render_previews_png(tmp_path: Path, template: Path) -> None:
    spec = write_spec(tmp_path, [{"layout": "Title Only", "title": "화재 신고"}])
    out = pf.build_deck(template, spec, tmp_path / "o.pptx")
    pdf, pngs = pf.render_previews(out, 40)
    assert pdf.stat().st_size > 0 and len(pngs) == 1 and pngs[0].suffix == ".png"
