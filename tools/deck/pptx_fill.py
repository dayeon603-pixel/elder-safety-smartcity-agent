"""Fill an official .pptx template from a JSON/YAML deck spec.

Usage:
    pptx_fill.py inspect template.pptx
    pptx_fill.py build template.pptx spec.json out.pptx [--no-render] [--dpi 60]
    pptx_fill.py default-template out.pptx

Spec format (JSON; YAML too if PyYAML is installed):
    {"slides": [{
        "layout": "Title Only" | 5,          # name (case-insensitive) or index
        "title": "...",
        "body": ["line 1", "line 2", "line 3"],   # max 3 lines
        "image": "photo.jpg",                # relative to the spec file
        "credit": "사진: ..., CC BY-SA 4.0, Wikimedia Commons",
        "notes": "speaker script",
        "placeholders": {"title": 0, "body": 1, "image": 13},  # optional idx overrides
        "image_box": [left, top, width, height]                 # optional, inches
    }]}
"""
from __future__ import annotations

import argparse
import json
import logging
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide, SlideLayout
from pptx.util import Emu, Inches, Pt

LOG = logging.getLogger("pptx_fill")

MAX_BODY_LINES = 3
WARN_LINE_CHARS = 20
CREDIT_FONT_PT = 9
CREDIT_RGB = RGBColor(0x80, 0x80, 0x80)
CREDIT_HEIGHT_IN = 0.3
CREDIT_BOTTOM_MARGIN_IN = 0.1
SIDE_MARGIN_RATIO = 0.05
GAP_IN = 0.15
SIDE_BODY_WIDTH_RATIO = 0.42
MIN_BELOW_HEIGHT_RATIO = 0.25
DEFAULT_PREVIEW_DPI = 60
SOFFICE_TIMEOUT_S = 240
PDFTOPPM_TIMEOUT_S = 120
PREVIEW_FONT = "AppleGothic"
PREVIEW_FONT_FILE = Path("/System/Library/Fonts/Supplemental/AppleGothic.ttf")
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
BODY_TYPES = {
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.SUBTITLE,
    PP_PLACEHOLDER.VERTICAL_BODY,
    PP_PLACEHOLDER.VERTICAL_OBJECT,
}
PICTURE_TYPES = {PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.BITMAP}


class SpecError(ValueError):
    """Invalid deck spec or template mapping."""


class RenderError(RuntimeError):
    """PDF/PNG conversion failed."""


@dataclass(frozen=True)
class Box:
    """Rectangle in EMU."""

    left: int
    top: int
    width: int
    height: int

    @property
    def bottom(self) -> int:
        return self.top + self.height


def _ph_type(shape: Any) -> PP_PLACEHOLDER | None:
    return shape.placeholder_format.type if shape.is_placeholder else None


def inspect_template(path: Path) -> list[str]:
    """Describe layouts and placeholders (idx, type, name, geometry in inches)."""
    prs = Presentation(str(path))
    out = [f"slide size: {Emu(prs.slide_width).inches:.2f} x {Emu(prs.slide_height).inches:.2f} in"]
    for i, layout in enumerate(prs.slide_layouts):
        out.append(f"[{i}] layout '{layout.name}'")
        for ph in layout.placeholders:
            pf = ph.placeholder_format
            geo = ""
            if ph.width is not None:
                geo = (
                    f" @ ({Emu(ph.left).inches:.2f},{Emu(ph.top).inches:.2f}) "
                    f"{Emu(ph.width).inches:.2f}x{Emu(ph.height).inches:.2f}in"
                )
            out.append(f"      idx={pf.idx:<3} type={pf.type.name if pf.type else '?':<14} name='{ph.name}'{geo}")
    return out


def resolve_layout(prs: PresentationType, ref: str | int) -> SlideLayout:
    """Find a layout by index or case-insensitive name."""
    layouts = list(prs.slide_layouts)
    if isinstance(ref, int) or (isinstance(ref, str) and ref.strip().isdigit()):
        idx = int(ref)
        if not 0 <= idx < len(layouts):
            raise SpecError(f"layout index {idx} out of range 0..{len(layouts) - 1}")
        return layouts[idx]
    for layout in layouts:
        if layout.name.strip().lower() == str(ref).strip().lower():
            return layout
    names = ", ".join(f"'{lo.name}'" for lo in layouts)
    raise SpecError(f"layout '{ref}' not found; available: {names}")


def load_spec(path: Path) -> dict[str, Any]:
    """Load spec from JSON or YAML and validate the basics."""
    text = path.read_text(encoding="utf-8")
    if path.suffix.lower() in {".yaml", ".yml"}:
        try:
            import yaml
        except ImportError as exc:
            raise SpecError("YAML spec needs PyYAML; use JSON instead") from exc
        data = yaml.safe_load(text)
    else:
        try:
            data = json.loads(text)
        except json.JSONDecodeError as exc:
            raise SpecError(f"invalid JSON in {path}: {exc}") from exc
    if not isinstance(data, dict) or not isinstance(data.get("slides"), list) or not data["slides"]:
        raise SpecError("spec must be an object with a non-empty 'slides' list")
    for n, s in enumerate(data["slides"], 1):
        if "layout" not in s:
            raise SpecError(f"slide {n}: 'layout' is required")
        body = s.get("body", [])
        if not isinstance(body, list) or len(body) > MAX_BODY_LINES:
            raise SpecError(f"slide {n}: 'body' must be a list of at most {MAX_BODY_LINES} lines")
        for line in body:
            if len(str(line)) > WARN_LINE_CHARS:
                LOG.warning("slide %d: body line over %d chars: %s", n, WARN_LINE_CHARS, line)
    return data


def contain_fit(img_w: int, img_h: int, box: Box) -> Box:
    """Largest box of the image's aspect ratio inside `box`, centered (no distortion)."""
    if img_w <= 0 or img_h <= 0 or box.width <= 0 or box.height <= 0:
        raise SpecError("image and target box must have positive size")
    scale = min(box.width / img_w, box.height / img_h)
    w, h = int(img_w * scale), int(img_h * scale)
    return Box(box.left + (box.width - w) // 2, box.top + (box.height - h) // 2, w, h)


def _find_by_idx(slide: Slide, idx: int) -> Any:
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    raise SpecError(f"no placeholder with idx={idx} on layout '{slide.slide_layout.name}'")


def _pick(slide: Slide, kinds: set[PP_PLACEHOLDER], override: int | None) -> list[Any]:
    if override is not None:
        return [_find_by_idx(slide, override)]
    found = [ph for ph in slide.placeholders if _ph_type(ph) in kinds]
    return sorted(found, key=lambda p: p.placeholder_format.idx)


def _remove(shape: Any) -> None:
    shape._element.getparent().remove(shape._element)


def _set_lines(shape: Any, lines: list[str]) -> None:
    tf = shape.text_frame
    tf.text = str(lines[0])
    for line in lines[1:]:
        tf.add_paragraph().text = str(line)


def _box_of(shape: Any) -> Box:
    return Box(int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def add_credit(prs: PresentationType, slide: Slide, credit: str) -> None:
    """Small gray credit line at the slide bottom."""
    left = int(prs.slide_width * SIDE_MARGIN_RATIO)
    width = int(prs.slide_width - 2 * left)
    height = Inches(CREDIT_HEIGHT_IN)
    top = int(prs.slide_height - height - Inches(CREDIT_BOTTOM_MARGIN_IN))
    tb = slide.shapes.add_textbox(left, top, width, height)
    tb.name = "Credit"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.text = credit
    run = tf.paragraphs[0].runs[0]
    run.font.size = Pt(CREDIT_FONT_PT)
    run.font.color.rgb = CREDIT_RGB


def _place_image(
    prs: PresentationType,
    slide: Slide,
    spec: dict[str, Any],
    image_path: Path,
    title_ph: Any | None,
    body_ph: Any | None,
) -> None:
    overrides: dict[str, int] = spec.get("placeholders", {})
    pic_phs = _pick(slide, PICTURE_TYPES, overrides.get("image"))
    pic_phs = [p for p in pic_phs if _ph_type(p) in PICTURE_TYPES or "image" in overrides]
    slide_w, slide_h = int(prs.slide_width), int(prs.slide_height)
    margin = int(slide_w * SIDE_MARGIN_RATIO)
    credit_top = slide_h - int(Inches(CREDIT_HEIGHT_IN + CREDIT_BOTTOM_MARGIN_IN))
    gap = int(Inches(GAP_IN))

    if "image_box" in spec:
        l, t, w, h = (Inches(float(v)) for v in spec["image_box"])
        box = Box(int(l), int(t), int(w), int(h))
    elif pic_phs:
        box = _box_of(pic_phs[0])
    else:
        top = (_box_of(title_ph).bottom + gap) if title_ph is not None else margin
        if body_ph is not None and body_ph.has_text_frame and body_ph.text_frame.text:
            # side-by-side: shrink body to the left, image takes the right
            body_w = int((slide_w - 2 * margin) * SIDE_BODY_WIDTH_RATIO)
            body_ph.left, body_ph.top = margin, top
            body_ph.width, body_ph.height = body_w, credit_top - gap - top
            box = Box(margin + body_w + gap, top, slide_w - margin - (margin + body_w + gap), credit_top - gap - top)
        else:
            box = Box(margin, top, slide_w - 2 * margin, credit_top - gap - top)
            if box.height < slide_h * MIN_BELOW_HEIGHT_RATIO:
                LOG.warning("little room for image on layout '%s'; use image_box", slide.slide_layout.name)
    for p in pic_phs:
        _remove(p)
    with Image.open(image_path) as im:
        iw, ih = im.size
    fit = contain_fit(iw, ih, box)
    pic = slide.shapes.add_picture(str(image_path), fit.left, fit.top, fit.width, fit.height)
    pic.name = "Image"


def build_deck(template: Path, spec_path: Path, out: Path) -> Path:
    """Create `out` from template + spec."""
    spec = load_spec(spec_path)
    prs = Presentation(str(template))
    for n, s in enumerate(spec["slides"], 1):
        slide = prs.slides.add_slide(resolve_layout(prs, s["layout"]))
        overrides: dict[str, int] = s.get("placeholders", {})
        titles = _pick(slide, TITLE_TYPES, overrides.get("title"))
        bodies = _pick(slide, BODY_TYPES, overrides.get("body"))
        title_ph = titles[0] if titles else None
        body_ph = bodies[0] if bodies else None
        if s.get("title"):
            if title_ph is None:
                raise SpecError(f"slide {n}: layout '{slide.slide_layout.name}' has no title placeholder")
            title_ph.text_frame.text = str(s["title"])
        lines = [str(x) for x in s.get("body", [])]
        if lines:
            if body_ph is None:
                raise SpecError(f"slide {n}: layout '{slide.slide_layout.name}' has no body placeholder")
            _set_lines(body_ph, lines)
        if s.get("image"):
            img = Path(s["image"])
            img = img if img.is_absolute() else spec_path.parent / img
            if not img.is_file():
                raise SpecError(f"slide {n}: image not found: {img}")
            _place_image(prs, slide, s, img, title_ph, body_ph if lines else None)
        if s.get("credit"):
            add_credit(prs, slide, str(s["credit"]))
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])
        # drop untouched empty placeholders so no prompt text shows in edit view
        for ph in list(slide.placeholders):
            if ph.has_text_frame and not ph.text_frame.text.strip() and _ph_type(ph) not in PICTURE_TYPES:
                _remove(ph)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    LOG.info("wrote %s (%d slides)", out, len(spec["slides"]))
    return out


def _run(cmd: list[str], timeout: int) -> None:
    try:
        subprocess.run(cmd, check=True, capture_output=True, timeout=timeout)
    except FileNotFoundError as exc:
        raise RenderError(f"{cmd[0]} not found on PATH") from exc
    except subprocess.TimeoutExpired as exc:
        raise RenderError(f"{cmd[0]} timed out after {timeout}s") from exc
    except subprocess.CalledProcessError as exc:
        raise RenderError(f"{cmd[0]} failed: {exc.stderr.decode(errors='replace')[:500]}") from exc


def make_preview_copy(src: Path, dst: Path, font: str) -> None:
    """Save a copy of `src` with `font` forced on every run (latin + East Asian).

    LibreOffice has no Hangul fallback for fonts such as Calibri, so previews of
    Korean text would otherwise render blank. The delivered .pptx is untouched.
    """
    prs = Presentation(str(src))

    def force(text_frame: Any) -> None:
        for para in text_frame.paragraphs:
            for run in para.runs:
                run.font.name = font
                rpr = run._r.get_or_add_rPr()
                for tag in ("ea", "cs"):
                    el = rpr.find(f"{{{A_NS}}}{tag}")
                    if el is None:
                        el = etree.SubElement(rpr, f"{{{A_NS}}}{tag}")
                    el.set("typeface", font)

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                force(shape.text_frame)
    prs.save(str(dst))


def render_previews(
    pptx: Path, dpi: int, font: str = PREVIEW_FONT, font_file: Path | None = PREVIEW_FONT_FILE
) -> tuple[Path, list[Path]]:
    """Convert to PDF (soffice) then PNG pages (pdftoppm). Returns (pdf, pngs).

    Rendering uses a temporary copy with `font` forced (see make_preview_copy).
    """
    with tempfile.TemporaryDirectory(prefix="deck_preview_") as work:
        work_dir = Path(work)
        tmp_pptx = work_dir / pptx.name
        make_preview_copy(pptx, tmp_pptx, font)
        tmp_pdf, tmp_pngs = _render(tmp_pptx, dpi, font_file)
        pdf = pptx.with_suffix(".pdf")
        for old in pptx.parent.glob(f"{pptx.stem}-*.png"):
            old.unlink()
        shutil.copy2(tmp_pdf, pdf)
        pngs = [Path(shutil.copy2(p, pptx.parent / p.name)) for p in tmp_pngs]
    return pdf, sorted(pngs)


def _render(pptx: Path, dpi: int, font_file: Path | None) -> tuple[Path, list[Path]]:
    pdf = pptx.with_suffix(".pdf")
    with tempfile.TemporaryDirectory(prefix="soffice_profile_") as profile:
        if font_file is not None and font_file.is_file():
            # headless soffice does not always see system fonts; the profile's
            # user/fonts directory is always scanned
            fonts_dir = Path(profile) / "user" / "fonts"
            fonts_dir.mkdir(parents=True)
            shutil.copyfile(font_file, fonts_dir / font_file.name)
        _run(
            [
                "soffice",
                f"-env:UserInstallation={Path(profile).as_uri()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(pptx.parent),
                str(pptx),
            ],
            SOFFICE_TIMEOUT_S,
        )
    if not pdf.is_file():
        raise RenderError(f"soffice produced no PDF at {pdf}")
    prefix = pptx.with_suffix("")
    _run(["pdftoppm", "-r", str(dpi), "-png", str(pdf), str(prefix)], PDFTOPPM_TIMEOUT_S)
    pngs = sorted(pptx.parent.glob(f"{prefix.name}-*.png"))
    return pdf, pngs


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    sub = p.add_subparsers(dest="cmd", required=True)
    i = sub.add_parser("inspect", help="list layouts and placeholders")
    i.add_argument("template", type=Path)
    b = sub.add_parser("build", help="fill template from spec")
    b.add_argument("template", type=Path)
    b.add_argument("spec", type=Path)
    b.add_argument("out", type=Path)
    b.add_argument("--no-render", action="store_true", help="skip PDF/PNG previews")
    b.add_argument("--dpi", type=int, default=DEFAULT_PREVIEW_DPI)
    b.add_argument("--preview-font", default=PREVIEW_FONT, help="font forced in preview copy only")
    d = sub.add_parser("default-template", help="save python-pptx default template (for testing)")
    d.add_argument("out", type=Path)
    return p


def main(argv: list[str] | None = None) -> int:
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    args = build_parser().parse_args(argv)
    try:
        if args.cmd == "inspect":
            for line in inspect_template(args.template):
                LOG.info("%s", line)
        elif args.cmd == "default-template":
            args.out.parent.mkdir(parents=True, exist_ok=True)
            Presentation().save(str(args.out))
            LOG.info("wrote %s", args.out)
        else:
            build_deck(args.template, args.spec, args.out)
            if not args.no_render:
                pdf, pngs = render_previews(args.out, args.dpi, args.preview_font)
                LOG.info("pdf: %s", pdf)
                for png in pngs:
                    LOG.info("png: %s", png)
        return 0
    except (SpecError, RenderError, OSError, KeyError) as exc:
        LOG.error("%s: %s", type(exc).__name__, exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
